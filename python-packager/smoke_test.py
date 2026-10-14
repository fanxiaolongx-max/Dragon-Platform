# -*- coding: utf-8 -*-
"""Exercise the portable program in an isolated path with Chinese and spaces.

The frozen EXE is tested by default. --source runs the same scenarios through
the source launcher, which is useful on macOS/Linux before the Windows CI run.
"""

import argparse
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import struct


PROJECT_DIR = Path(__file__).resolve().parent
EXCEL_PROBE = '''# -*- coding: utf-8 -*-
import glob
from pathlib import Path
import pandas as pd
import xlrd
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pandas.io.excel import _openpyxl, _xlrd

data = Path(__file__).resolve().parent / "data"
data.mkdir(exist_ok=True)
source = data / "中文输入 测试.xlsx"
workbook = Workbook()
sheet = workbook.active
sheet.title = "Task List"
sheet.append(["风险ID", "说明", "数值"])
sheet.append(["R-001", "中文正常", 42])
sheet["A1"].font = Font(bold=True)
workbook.save(source)
frame = pd.read_excel(source, sheet_name="Task List", engine="openpyxl")
assert frame.iloc[0].tolist() == ["R-001", "中文正常", 42], frame
legacy_frame = pd.read_excel(data / "中文旧格式 测试.xls", sheet_name="Task List", engine="xlrd")
assert legacy_frame.iloc[0].tolist() == ["R-001", "中文正常", 42], legacy_frame
target = data / "中文输出 报告.xlsx"
frame.to_excel(target, index=False, engine="openpyxl")
assert load_workbook(target).active["B2"].value == "中文正常"
assert len(glob.glob(str(data / "*.xlsx"))) == 2
assert xlrd.__version__
if sys.platform == 'win32':
    import ctypes
    assert ctypes.windll.kernel32.GetConsoleWindow() == 0, 'Worker must run without a console window'
print("EXCEL_SMOKE_OK 中文读写成功", flush=True)
'''


def run_core(command, cwd, input_text, markers, timeout):
    env = os.environ.copy()
    env.update({"PYTHONUTF8": "1", "PYTHONIOENCODING": "utf-8"})
    # Do not let an installed module or caller's Python configuration mask a
    # missing dependency in the frozen bundle.
    env.pop("PYTHONPATH", None)
    env.pop("PYTHONHOME", None)
    result = subprocess.run(
        command, cwd=cwd, input=input_text, encoding="utf-8", errors="strict",
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT, env=env, timeout=timeout,
        creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
    )
    if result.returncode != 0:
        raise RuntimeError(f"Core exited with code {result.returncode}:\n{result.stdout[-6000:]}")
    for marker in markers:
        if marker not in result.stdout:
            raise RuntimeError(f"Missing output {marker!r}:\n{result.stdout[-6000:]}")
    if "Traceback (most recent call last)" in result.stdout:
        raise RuntimeError(f"Unexpected traceback:\n{result.stdout[-6000:]}")
    return result.stdout


def _windows_subsystem(executable):
    data = Path(executable).read_bytes()
    header = struct.unpack_from('<I', data, 0x3c)[0]
    if data[header:header + 4] != b'PE\0\0':
        raise RuntimeError('Invalid Windows executable')
    return struct.unpack_from('<H', data, header + 4 + 20 + 68)[0]


def smoke_test(dist_dir=None, exe_name="PR-Audit-Worker.exe", source=False, timeout=60, gui_exe="PR小助手.exe"):
    with tempfile.TemporaryDirectory(prefix="PR 冒烟测试 ") as temp_dir:
        portable_dir = Path(temp_dir) / "中文目录 with spaces" / "PR 报告工具"
        if source:
            portable_dir.mkdir(parents=True)
            for filename in ("universal_shell.py", "core_script.py", "version.json", "updater_config.json"):
                shutil.copy2(PROJECT_DIR / filename, portable_dir / filename)
            command = [sys.executable, str(portable_dir / "universal_shell.py"), "--run-core"]
            gui_command = [sys.executable, str(portable_dir / 'universal_shell.py')]
        else:
            if dist_dir is None:
                raise ValueError("--dist-dir is required for the EXE smoke test")
            shutil.copytree(Path(dist_dir), portable_dir)
            executable = portable_dir / exe_name
            if not executable.is_file():
                raise FileNotFoundError(f"Missing portable executable: {exe_name}")
            command = [str(executable), "--run-core"]
            gui_command = [str(portable_dir / gui_exe)]
            if sys.platform == 'win32':
                if _windows_subsystem(gui_command[0]) != 2 or _windows_subsystem(executable) != 3:
                    raise RuntimeError('Expected a windowed GUI and console-enabled background worker')
        gui_result = portable_dir / 'gui-smoke.json'
        subprocess.run(gui_command + ['--gui-smoke-result', str(gui_result)], cwd=portable_dir,
                       env={**os.environ, 'PYTHONUTF8': '1', 'PYTHONIOENCODING': 'utf-8'},
                       timeout=timeout, check=True,
                       creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == 'win32' else 0)
        if not json.loads(gui_result.read_text(encoding='utf-8')).get('passed'):
            raise RuntimeError(f'GUI smoke failed: {gui_result.read_text(encoding="utf-8")}')
        print('PASS: windowed GUI, background worker and inline terminal input', flush=True)
        run_core(command, portable_dir, "0\n", ("主菜单", "请选择", "再见"), timeout)
        print("PASS: real core menu, UTF-8 input/output and clean exit", flush=True)
        probe_bytes = EXCEL_PROBE.encode("utf-8")
        (portable_dir / "core_script.py").write_bytes(probe_bytes)
        (portable_dir / "data").mkdir(exist_ok=True)
        # Synthetic BIFF8 fixture, generated with xlwt from just two fake rows.
        # xlwt is not needed by the application or on the CI build machine.
        shutil.copy2(PROJECT_DIR / "tests" / "fixtures" / "excel_smoke.xls",
                     portable_dir / "data" / "中文旧格式 测试.xls")
        # Keep the temporary core/metadata pair consistent for any launcher
        # integrity checks. The real portable files remain untouched.
        metadata_path = portable_dir / "version.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8-sig"))
        metadata["sha256"] = hashlib.sha256(probe_bytes).hexdigest()
        metadata_path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")
        run_core(command, portable_dir, "", ("EXCEL_SMOKE_OK 中文读写成功",), timeout)
        print("PASS: glob, pandas/openpyxl .xlsx round trip and xlrd .xls read", flush=True)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dist-dir", type=Path)
    parser.add_argument("--exe", default="PR-Audit-Worker.exe")
    parser.add_argument('--gui-exe', default='PR小助手.exe')
    parser.add_argument("--source", action="store_true")
    parser.add_argument("--timeout", type=float, default=60)
    args = parser.parse_args(argv)
    try:
        smoke_test(args.dist_dir, args.exe, args.source, args.timeout, args.gui_exe)
        return 0
    except (OSError, ValueError, RuntimeError, UnicodeError, subprocess.TimeoutExpired) as error:
        print(f"Smoke test failed: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
